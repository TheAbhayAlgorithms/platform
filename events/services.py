import io
import csv
from decimal import Decimal
from typing import Dict, Any, Optional
from django.db import transaction
from django.utils import timezone
from django.core.exceptions import ValidationError
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from members.models import Member, MemberStatus
from notifications.models import NotificationTemplate
from notifications.services import send_whatsapp_notification
from events.models import (
    Event,
    EventRegistration,
    EventStatus,
    RegistrationStatus,
    EventPaymentStatus,
)


def register_member_for_event(
    event: Event,
    member: Member,
    num_family_members: int = 0,
    num_guests: int = 0,
    notes: str = '',
    registered_by: Optional[Any] = None
) -> EventRegistration:
    """
    Registers a member and accompanying attendees for an institute festival or event.
    Enforces capacity limits, deadlines, guest quotas, and sends WhatsApp confirmation.
    """
    if member.is_removed:
        raise ValidationError("Cannot register a removed member for institute events.")

    if member.status == MemberStatus.SUSPENDED:
        raise ValidationError("Membership is suspended. Event privileges are withheld.")

    if not event.is_active or event.status != EventStatus.PUBLISHED:
        raise ValidationError(f"Event '{event.title}' is not currently open for registrations.")

    now = timezone.now()
    if now >= event.end_date_time:
        raise ValidationError("This event has already concluded.")

    if event.registration_deadline and now > event.registration_deadline:
        raise ValidationError(
            f"The registration deadline for this event closed on {event.registration_deadline:%d-%m-%Y %H:%M}."
        )

    if num_guests < 0 or num_family_members < 0:
        raise ValidationError("Attendee counts cannot be negative.")

    if num_guests > event.max_guests_per_member:
        raise ValidationError(
            f"You may bring at most {event.max_guests_per_member} outside guests for this event."
        )

    with transaction.atomic():
        # Check for existing active registration
        existing = EventRegistration.objects.filter(
            event=event,
            member=member,
            status=RegistrationStatus.CONFIRMED
        ).first()

        if existing:
            raise ValidationError(
                f"{member.name} ({member.membership_id}) is already registered for this event."
            )

        # Capacity check
        requested_headcount = 1 + num_family_members + num_guests
        if event.capacity > 0:
            current_headcount = event.registered_headcount
            if (current_headcount + requested_headcount) > event.capacity:
                available = max(0, event.capacity - current_headcount)
                raise ValidationError(
                    f"Insufficient capacity remaining. Available seats: {available}, requested: {requested_headcount}."
                )

        # Fee computation
        fee = event.member_entry_fee + (Decimal(num_guests) * event.guest_entry_fee)
        payment_status = EventPaymentStatus.FREE if fee <= Decimal('0.00') else EventPaymentStatus.PENDING

        registration = EventRegistration.objects.create(
            event=event,
            member=member,
            num_family_members=num_family_members,
            num_guests=num_guests,
            calculated_fee=fee,
            payment_status=payment_status,
            status=RegistrationStatus.CONFIRMED,
            notes=notes.strip() if notes else '',
            registered_by=registered_by
        )

        # Automated WhatsApp confirmation
        if member.whatsapp_opt_in:
            confirm_msg = (
                f"Namaste {member.name},\n\n"
                f"Your RSVP for '{event.title}' has been successfully confirmed!\n\n"
                f"Date & Time: {event.start_date_time:%d-%m-%Y at %I:%M %p}\n"
                f"Venue: {event.venue}\n"
                f"Total Attendees: {registration.total_headcount} "
                f"(Member: 1, Family: {num_family_members}, Guests: {num_guests})\n"
            )
            if fee > Decimal('0.00'):
                confirm_msg += f"Entry Fee: ₹{fee:,.2f} ({registration.get_payment_status_display()})\n"
            else:
                confirm_msg += "Entry: Complimentary for Member Family\n"

            confirm_msg += (
                f"\nPlease display your Membership ID ({member.membership_id}) at the entrance desk for check-in.\n\n"
                f"Helpline: 7039145304"
            )

            send_whatsapp_notification(
                member=member,
                template_type=NotificationTemplate.CUSTOM,
                custom_message=confirm_msg
            )

        return registration


def cancel_event_registration(registration: EventRegistration, reason: str = '') -> EventRegistration:
    """Cancels an active event registration, releasing seats back to event capacity."""
    registration.status = RegistrationStatus.CANCELLED
    if reason:
        registration.notes = f"{registration.notes}\n[Cancelled: {reason.strip()}]".strip()
    registration.save(update_fields=['status', 'notes'])
    return registration


def toggle_event_check_in(registration: EventRegistration) -> bool:
    """Toggles venue arrival check-in status for an attendee."""
    if registration.checked_in:
        registration.checked_in = False
        registration.checked_in_at = None
    else:
        registration.checked_in = True
        registration.checked_in_at = timezone.now()

    registration.save(update_fields=['checked_in', 'checked_in_at'])
    return registration.checked_in


def get_event_headcount_stats(event: Event) -> Dict[str, Any]:
    """Computes comprehensive registration and venue check-in statistics for an event."""
    confirmed_regs = event.registrations.filter(status=RegistrationStatus.CONFIRMED).select_related('member')

    total_registrations = confirmed_regs.count()
    member_count = total_registrations
    family_count = sum(r.num_family_members for r in confirmed_regs)
    guest_count = sum(r.num_guests for r in confirmed_regs)
    total_headcount = member_count + family_count + guest_count

    checked_in_regs = confirmed_regs.filter(checked_in=True)
    checked_in_members = checked_in_regs.count()
    checked_in_family = sum(r.num_family_members for r in checked_in_regs)
    checked_in_guests = sum(r.num_guests for r in checked_in_regs)
    total_checked_in = checked_in_members + checked_in_family + checked_in_guests

    total_fees = sum((r.calculated_fee for r in confirmed_regs), Decimal('0.00'))

    return {
        'total_registrations': total_registrations,
        'member_count': member_count,
        'family_count': family_count,
        'guest_count': guest_count,
        'total_headcount': total_headcount,
        'checked_in_members': checked_in_members,
        'checked_in_family': checked_in_family,
        'checked_in_guests': checked_in_guests,
        'total_checked_in': total_checked_in,
        'total_fees': total_fees,
        'capacity': event.capacity,
        'remaining_capacity': event.remaining_capacity,
        'is_full': event.is_full,
    }


def export_event_attendees_excel(event: Event) -> bytes:
    """Exports attendee roster to Microsoft Excel (.xlsx) format using openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendees"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Membership ID", "Member Name", "Phone", "Department",
        "Family", "Guests", "Total Headcount", "Fee (₹)",
        "Payment Status", "Checked In", "Notes", "Registration Time"
    ]

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"RAILWAY EMPLOYEES' INSTITUTE - {event.title.upper()} ATTENDEE ROSTER"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Subtitle
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = f"Venue: {event.venue} | Date: {event.start_date_time:%d-%m-%Y %H:%M} | Exported: {timezone.now():%d-%m-%Y %H:%M}"
    sub_cell.font = Font(name="Calibri", size=9, italic=True, color="4B5563")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Table Header Row
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.row_dimensions[4].height = 22
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # Data Rows
    regs = event.registrations.filter(status=RegistrationStatus.CONFIRMED).select_related('member').order_by('registered_at')
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    current_row = 5

    for r_idx, reg in enumerate(regs):
        ws.row_dimensions[current_row].height = 19
        fill_to_use = alt_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)

        row_vals = [
            reg.member.membership_id,
            reg.member.name,
            reg.member.phone,
            reg.member.get_department_display(),
            reg.num_family_members,
            reg.num_guests,
            reg.total_headcount,
            float(reg.calculated_fee),
            reg.get_payment_status_display(),
            "Yes" if reg.checked_in else "No",
            reg.notes or "-",
            reg.registered_at.strftime("%d-%m-%Y %H:%M")
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            if fill_to_use.fill_type:
                c.fill = fill_to_use
            if isinstance(val, (int, float, Decimal)):
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2]:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_event_attendees_csv(event: Event) -> str:
    """Exports attendee roster to standard CSV format."""
    output = io.StringIO()
    writer = csv.writer(output)
    headers = [
        "Membership ID", "Member Name", "Phone", "Department",
        "Family", "Guests", "Total Headcount", "Fee",
        "Payment Status", "Checked In", "Notes", "Registration Time"
    ]
    writer.writerow(headers)

    regs = event.registrations.filter(status=RegistrationStatus.CONFIRMED).select_related('member').order_by('registered_at')
    for reg in regs:
        writer.writerow([
            reg.member.membership_id,
            reg.member.name,
            reg.member.phone,
            reg.member.get_department_display(),
            reg.num_family_members,
            reg.num_guests,
            reg.total_headcount,
            reg.calculated_fee,
            reg.get_payment_status_display(),
            "Yes" if reg.checked_in else "No",
            reg.notes or "",
            reg.registered_at.strftime("%d-%m-%Y %H:%M")
        ])

    return output.getvalue()

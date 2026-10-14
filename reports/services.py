import io
import csv
from datetime import date, timedelta
from decimal import Decimal
from typing import Dict, Any, List, Optional
from django.db import models
from django.utils import timezone
from django.conf import settings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from members.models import Member, MemberStatus, DepartmentChoices
from payments.models import Payment, PaymentStatus, PaymentMode, PaymentAdjustment
from payments.services import calculate_member_dues
from plans.models import Plan, MemberPlanAssignment


def get_daily_collection_data(target_date: date) -> Dict[str, Any]:
    """
    Gathers detailed financial collections for a single calendar day.
    Includes payments list, total collected, transaction count, and payment mode breakdown.
    """
    payments = Payment.objects.filter(
        payment_date=target_date,
        status=PaymentStatus.COMPLETED,
        is_cancelled=False
    ).select_related('member', 'plan_assignment__plan', 'collected_by').order_by('-created_at')


    total_amount = Decimal('0.00')
    mode_counts = {mode.value: {'label': mode.label, 'count': 0, 'amount': Decimal('0.00')} for mode in PaymentMode}

    for p in payments:
        total_amount += p.amount
        if p.payment_mode in mode_counts:
            mode_counts[p.payment_mode]['count'] += 1
            mode_counts[p.payment_mode]['amount'] += p.amount

    # Also check adjustments / refunds on that date
    adjustments = PaymentAdjustment.objects.filter(
        created_at__date=target_date
    ).aggregate(total=models.Sum('amount'))['total'] or Decimal('0.00')

    net_collection = total_amount - adjustments

    return {
        'target_date': target_date,
        'payments': payments,
        'total_amount': total_amount,
        'total_count': payments.count(),
        'adjustments': adjustments,
        'net_collection': net_collection,
        'mode_breakdown': mode_counts,
    }


def get_revenue_summary_data(start_date: date, end_date: date) -> Dict[str, Any]:
    """
    Gathers aggregate revenue summary over a date range.
    Includes total collections, refunds/adjustments, net revenue,
    breakdown by payment mode, and breakdown by plan.
    """
    payments = Payment.objects.filter(
        payment_date__range=(start_date, end_date),
        status=PaymentStatus.COMPLETED,
        is_cancelled=False
    ).select_related('member', 'plan_assignment__plan')

    total_revenue = payments.aggregate(total=models.Sum('amount'))['total'] or Decimal('0.00')
    total_transactions = payments.count()

    # Payment Mode Breakdown
    mode_summary = {mode.value: {'label': mode.label, 'count': 0, 'amount': Decimal('0.00')} for mode in PaymentMode}
    for p in payments:
        if p.payment_mode in mode_summary:
            mode_summary[p.payment_mode]['count'] += 1
            mode_summary[p.payment_mode]['amount'] += p.amount

    # Plan Category Breakdown
    plan_summary = {}
    for p in payments:
        plan_name = p.plan_assignment.plan.name if (p.plan_assignment and p.plan_assignment.plan) else "General / Custom Fee"
        if plan_name not in plan_summary:
            plan_summary[plan_name] = {'count': 0, 'amount': Decimal('0.00')}
        plan_summary[plan_name]['count'] += 1
        plan_summary[plan_name]['amount'] += p.amount

    # Adjustments in period
    adjustments = PaymentAdjustment.objects.filter(
        created_at__date__range=(start_date, end_date)
    ).aggregate(total=models.Sum('amount'))['total'] or Decimal('0.00')

    net_revenue = total_revenue - adjustments

    return {
        'start_date': start_date,
        'end_date': end_date,
        'total_revenue': total_revenue,
        'total_transactions': total_transactions,
        'adjustments': adjustments,
        'net_revenue': net_revenue,
        'mode_summary': mode_summary,
        'plan_summary': plan_summary,
        'payments': payments[:100],  # preview capped at 100 for display
    }


def get_defaulters_data(department: Optional[str] = None) -> Dict[str, Any]:
    """
    Identifies all members with outstanding dues or expired status.
    Calculates overdue balance, active plan, days overdue, and total institute deficit.
    """
    today = timezone.now().date()
    members_qs = Member.objects.active_records().prefetch_related(
        'plan_assignments__plan',
        'payments'
    )

    if department:
        members_qs = members_qs.filter(department=department)

    defaulters = []
    total_outstanding = Decimal('0.00')

    for member in members_qs:
        dues_info = calculate_member_dues(member)
        balance = dues_info['balance_due']

        # Include if has positive balance or status is DUE or EXPIRED
        if balance > Decimal('0.00') or member.status in [MemberStatus.DUE, MemberStatus.EXPIRED]:
            current_assignment = dues_info.get('plan_assignment')
            end_date = current_assignment.end_date if current_assignment else None
            days_overdue = (today - end_date).days if end_date and today > end_date else 0

            defaulters.append({
                'member': member,
                'membership_id': member.membership_id,
                'name': member.name,
                'phone': member.phone,
                'department': member.get_department_display(),
                'status': member.get_status_display(),
                'plan_name': dues_info['plan_name'],
                'end_date': end_date,
                'days_overdue': max(0, days_overdue),
                'total_due': dues_info['total_due'],
                'total_paid': dues_info['total_paid'],
                'balance_due': balance,
            })
            total_outstanding += balance

    # Sort by balance_due descending
    defaulters.sort(key=lambda x: x['balance_due'], reverse=True)

    return {
        'defaulters': defaulters,
        'total_count': len(defaulters),
        'total_outstanding': total_outstanding,
        'department_filter': department,
        'departments': DepartmentChoices.choices,
    }


def get_member_demographics_data() -> Dict[str, Any]:
    """
    Gathers member demographic distributions by department, status,
    and plan assignment coverage.
    """
    active_members = Member.objects.active_records()
    total_members = active_members.count()

    # Department breakdown
    dept_counts = {}
    for dept_code, dept_label in DepartmentChoices.choices:
        count = active_members.filter(department=dept_code).count()
        if count > 0:
            dept_counts[dept_label] = count

    # Status breakdown
    status_counts = {}
    for st_code, st_label in MemberStatus.choices:
        if st_code != MemberStatus.REMOVED:
            status_counts[st_label] = active_members.filter(status=st_code).count()

    # Plan coverage
    with_plan_count = active_members.filter(plan_assignments__is_current=True).distinct().count()
    without_plan_count = total_members - with_plan_count

    # DPDP WhatsApp Opt-in distribution
    opt_in_count = active_members.filter(whatsapp_opt_in=True).count()
    opt_out_count = total_members - opt_in_count

    return {
        'total_members': total_members,
        'dept_counts': dept_counts,
        'status_counts': status_counts,
        'with_plan_count': with_plan_count,
        'without_plan_count': without_plan_count,
        'opt_in_count': opt_in_count,
        'opt_out_count': opt_out_count,
    }


# =========================================================================
# EXPORT GENERATORS: EXCEL (.xlsx), CSV, PDF
# =========================================================================

def export_to_excel(
    report_title: str,
    headers: List[str],
    rows: List[List[Any]],
    summary_metrics: Optional[Dict[str, Any]] = None
) -> bytes:
    """
    Generates a professionally styled Microsoft Excel spreadsheet (.xlsx) using openpyxl.
    Features branded header banner, auto column widths, borders, and bold summaries.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.views.sheetView[0].showGridLines = True

    # Title styling
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"RAILWAY EMPLOYEES' INSTITUTE - {report_title.upper()}"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Subtitle / Generation timestamp
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = f"Generated On: {timezone.now():%d-%m-%Y %H:%M} IST | Official Platform Audit Export"
    sub_cell.font = Font(name="Calibri", size=9, italic=True, color="6B7280")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    current_row = 4

    # Optional Summary Metrics Cards at the top
    if summary_metrics:
        for key, val in summary_metrics.items():
            ws.cell(row=current_row, column=1).value = f"{key}:"
            ws.cell(row=current_row, column=1).font = Font(name="Calibri", size=10, bold=True, color="374151")
            ws.cell(row=current_row, column=2).value = str(val)
            ws.cell(row=current_row, column=2).font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
            current_row += 1
        current_row += 1

    # Table Header Row
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.row_dimensions[current_row].height = 24
    for col_idx, header_text in enumerate(headers, start=1):
        c = ws.cell(row=current_row, column=col_idx, value=header_text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    current_row += 1

    # Data Rows
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    for r_idx, row_data in enumerate(rows):
        ws.row_dimensions[current_row].height = 20
        fill_to_use = alt_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)

        for col_idx, cell_value in enumerate(row_data, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=cell_value)
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            if fill_to_use.fill_type:
                c.fill = fill_to_use

            # Alignment logic
            if isinstance(cell_value, (int, float, Decimal)):
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title merged row
            if cell.row in [1, 2]:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_to_csv(headers: List[str], rows: List[List[Any]]) -> str:
    """
    Generates standard comma-separated values (CSV) string.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return output.getvalue()


def export_to_pdf(
    report_title: str,
    headers: List[str],
    rows: List[List[Any]],
    summary_metrics: Optional[Dict[str, Any]] = None,
    orientation: str = 'portrait'
) -> bytes:
    """
    Generates an official formatted PDF document using reportlab.
    Features header banner, institute styling, metadata, and structured tables.
    """
    buffer = io.BytesIO()
    pagesize = landscape(letter) if orientation == 'landscape' else letter

    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#4B5563'),
        alignment=1
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10
    )
    header_cell_style = ParagraphStyle(
        'HeaderCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )

    story = []

    # Institute Header
    story.append(Paragraph(f"RAILWAY EMPLOYEES' INSTITUTE &bull; {report_title.upper()}", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"Generated On: {timezone.now():%d-%m-%Y %H:%M} IST &bull; Platform Audit Report", sub_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#1E3A8A'), spaceBefore=2, spaceAfter=10))

    # Summary Metrics Table
    if summary_metrics:
        metric_data = []
        row_temp = []
        for k, v in summary_metrics.items():
            row_temp.extend([
                Paragraph(f"<b>{k}:</b>", cell_style),
                Paragraph(str(v), cell_style)
            ])
            if len(row_temp) >= 4:
                metric_data.append(row_temp)
                row_temp = []
        if row_temp:
            while len(row_temp) < 4:
                row_temp.append(Paragraph("", cell_style))
            metric_data.append(row_temp)

        metric_table = Table(metric_data, colWidths=[120, 150, 120, 150])
        metric_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F3F4F6')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(metric_table)
        story.append(Spacer(1, 12))

    # Main Data Table
    table_data = []
    # Header row
    table_data.append([Paragraph(h, header_cell_style) for h in headers])

    for row in rows:
        row_cells = []
        for val in row:
            text = str(val) if val is not None else ""
            row_cells.append(Paragraph(text, cell_style))
        table_data.append(row_cells)

    # Dynamic column widths
    num_cols = len(headers)
    page_width = pagesize[0] - 72  # 36 margin on each side
    col_width = page_width / num_cols

    main_table = Table(table_data, colWidths=[col_width] * num_cols, repeatRows=1)
    main_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
        ('TOPPADDING', (0, 0), (-1, 0), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0, 1), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
    ]))

    story.append(main_table)
    doc.build(story)
    return buffer.getvalue()
